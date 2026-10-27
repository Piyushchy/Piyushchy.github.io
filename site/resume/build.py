#!/usr/bin/env python3
"""
Builds the resume in every format from ONE data block (DATA below).

    python3 site/resume/build.py        (from the repo root, or anywhere)

Writes to site/public/resume/:
    index.html                      the /resume page on the site
    Piyush-Chaudhary-Resume.pdf     compiled from the .tex with pdflatex
    Piyush-Chaudhary-Resume.tex     LaTeX source (single column, no tables/images)
    Piyush-Chaudhary-Resume.docx    Word
    Piyush-Chaudhary-Resume.txt     plain text

Everything is deliberately plain so applicant tracking systems can read it:
one column, standard section names, real text (no images, icons or text boxes),
standard bullets, and a Unicode map in the PDF so copy/paste and parsers get clean text.
Edit DATA, run the script, then `pnpm build && pnpm publish:pages`.
"""
import html
import re
import shutil
import subprocess
import tempfile
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "public" / "resume"
BASE = "Piyush-Chaudhary-Resume"

DATA = {
    "name": "Piyush Chaudhary",
    "contact": [
        ("piyush.chaudhary@somaiya.edu", "mailto:piyush.chaudhary@somaiya.edu"),
        ("github.com/Piyushchy", "https://github.com/Piyushchy"),
        ("linkedin.com/in/piyush-chaudhary-9b5999187", "https://www.linkedin.com/in/piyush-chaudhary-9b5999187/"),
        ("piyushchy.github.io", "https://piyushchy.github.io"),
    ],
    "summary": (
        "Final-year B.Tech Computer Engineering student who builds document-processing, cloud and mobile software. "
        "Currently an intern at SBI Global IT Centre, previously an intern at Mazagon Dock Shipbuilders Limited."
    ),
    "education": [
        {
            "title": "Bachelor of Technology, Computer Engineering",
            "org": "K.J. Somaiya College of Engineering",
            "when": "Final year",
            "lines": ["Coursework: Natural Language Processing, Cloud Computing, AR/VR, Network and Information Security"],
        }
    ],
    "experience": [
        {
            "title": "Intern",
            "org": "State Bank of India, Global IT Centre",
            "when": "Present",
            "bullets": [
                "Building a production pipeline that extracts structured tables from scanned telecom audit PDFs into Excel, "
                "using a hybrid of computer vision (PyMuPDF, OpenCV) and a vision-language model.",
                "Compared three architectures on the same documents (Table Transformer, a pure vision-language model, "
                "and Docling) before choosing the hybrid design.",
                "Added a fast pre-scan that locates key header fields and the conclusion-table page range before the "
                "per-page pass, and switched to lazy page generators after eager rendering exhausted memory.",
                "Fixed a column-splicing bug that turned an 11-column table into 24 phantom columns, and added a guard "
                "that rejects implausible vision-language model cell text.",
                "Generated synthetic scanned PDFs with noise and skew for testing, and documented the system in a "
                "16-page technical report written in LaTeX.",
            ],
        },
        {
            "title": "Intern",
            "org": "Mazagon Dock Shipbuilders Limited",
            "when": "",
            "bullets": [],
        },
    ],
    "projects": [
        {
            "title": "The Tangled Path: Choice-Based Anti-Drug Game",
            "stack": "HTML, CSS, JavaScript",
            "link": ("sudo-choices.vercel.app", "https://sudo-choices.vercel.app/"),
            "bullets": [
                "1st runner up at a hackathon. Team leader and front-end developer on a team of four.",
                "A story game that follows one character through four life stages (ages 6, 15, 25 and 70), where each "
                "choice leads toward a good or a bad ending.",
                "Includes a built-in story teller and a full-screen mode. Deployed on Vercel.",
            ],
        },
        {
            "title": "MouseShare",
            "stack": "Python, Kotlin, Android",
            "link": None,
            "bullets": [
                "Shares a Windows PC's mouse and keyboard with an Android phone: a Python desktop app captures input and "
                "a Kotlin Android app receives and acts on it.",
                "Comparing transports (HID keyboard and mouse, WiFi) to cut input lag.",
            ],
        },
        {
            "title": "GIF Crop and Overlay Tool",
            "stack": "JavaScript, HTML, Canvas",
            "link": ("piyushchy.github.io/gif-tool", "https://piyushchy.github.io/gif-tool/"),
            "bullets": [
                "Browser tool that crops a GIF to a freehand or polygon shape and overlays one GIF on another with "
                "per-frame controls (blank, skip, hold, and shifting one track against the other).",
                "Wrote the GIF decoder and encoder from scratch (LZW and median-cut palette). Runs entirely in the "
                "browser, so nothing is uploaded.",
            ],
        },
        {
            "title": "Cloud-Native App on AWS",
            "stack": "Docker, AWS ECR, ECS Fargate, ALB, RDS MySQL, CodePipeline, SNS",
            "link": None,
            "bullets": [
                "Containerized a Node.js app and deployed it on ECS Fargate behind an Application Load Balancer, with "
                "the image stored in ECR.",
                "Connected to an RDS MySQL database through a bastion host over Session Manager, and added a "
                "CodePipeline CI/CD flow and an SNS alarm.",
            ],
        },
        {
            "title": "RNN vs LSTM Text Classification",
            "stack": "Python, PyTorch, Google Colab",
            "link": None,
            "bullets": [
                "Trained and compared a plain RNN and an LSTM on the same text classification task in a runnable "
                "Google Colab notebook.",
            ],
        },
    ],
    "skills": [
        ("Languages", "Python, Kotlin, JavaScript, HTML, CSS, LaTeX"),
        ("Computer vision and document AI", "OpenCV, PyMuPDF, EasyOCR, Table Transformer, Docling, Pillow, vision-language models"),
        ("Machine learning", "PyTorch, RNN, LSTM"),
        ("Cloud and infrastructure", "Docker, AWS (ECR, ECS Fargate, ALB, RDS MySQL, EC2, CodePipeline, SNS)"),
        ("Tools", "Git, GitHub, Google Colab, openpyxl, Excel automation"),
    ],
    "awards": ["1st runner up, hackathon, for The Tangled Path (team leader and front-end developer)"],
}


# ---------------------------------------------------------------- LaTeX
def tex_esc(s: str) -> str:
    s = s.replace("\\", r"\textbackslash{}")
    for a, b in (("&", r"\&"), ("%", r"\%"), ("$", r"\$"), ("#", r"\#"), ("_", r"\_"), ("{", r"\{"), ("}", r"\}")):
        s = s.replace(a, b)
    return s.replace("~", r"\textasciitilde{}").replace("^", r"\textasciicircum{}")


def build_tex(d) -> str:
    def entry_head(left_bold, left_rest, right):
        r = rf"\hfill {tex_esc(right)}" if right else ""
        rest = f", {tex_esc(left_rest)}" if left_rest else ""
        return rf"\noindent\textbf{{{tex_esc(left_bold)}}}{rest}{r}\par"

    def bullets(items):
        if not items:
            return ""
        return "\\begin{itemize}\n" + "\n".join(rf"  \item {tex_esc(b)}" for b in items) + "\n\\end{itemize}\n"

    L = []
    L.append(r"""\documentclass[10pt,letterpaper]{article}
% ATS-friendly: one column, standard fonts, no tables, no images, no text boxes.
\usepackage[T1]{fontenc}
\usepackage[utf8]{inputenc}
\usepackage[scaled=0.92]{helvet}  % Arial-style vector font (Type 1), the safest choice for parsers
\renewcommand{\familydefault}{\sfdefault}
\usepackage{cmap}                 % makes the PDF text copy/paste and parse cleanly
\usepackage[letterpaper,left=0.65in,right=0.65in,top=0.45in,bottom=0.45in]{geometry}
\usepackage{microtype}
\DisableLigatures{encoding = *, family = *}   % "fi" and "fl" stay two plain letters for parsers
\usepackage{enumitem}
\usepackage[hidelinks]{hyperref}
\usepackage{titlesec}
\input{glyphtounicode}
\pdfgentounicode=1
\pagestyle{empty}
\setlength{\parindent}{0pt}
\setlength{\parskip}{0pt}
\raggedright
\titleformat{\section}{\large\bfseries}{}{0em}{}[\titlerule]
\titlespacing*{\section}{0pt}{6pt}{3pt}
\setlist[itemize]{leftmargin=1.2em,itemsep=0pt,topsep=1pt,parsep=0pt}
\begin{document}
""")
    L.append(rf"\begin{{center}}{{\LARGE\bfseries {tex_esc(d['name'])}}}\\[3pt]")
    L.append(" \\textbar{} ".join(rf"\href{{{u}}}{{{tex_esc(t)}}}" for t, u in d["contact"]) + "\n\\end{center}\n")
    L.append("\\section*{SUMMARY}\n" + tex_esc(d["summary"]) + "\n")
    L.append("\\section*{EDUCATION}")
    for e in d["education"]:
        L.append(entry_head(e["title"], e["org"], e["when"]))
        for ln in e["lines"]:
            L.append(tex_esc(ln) + "\\par")
    L.append("\n\\section*{EXPERIENCE}")
    for i, e in enumerate(d["experience"]):
        if i:
            L.append("\\vspace{1pt}")
        L.append(entry_head(e["title"], e["org"], e["when"]))
        L.append(bullets(e["bullets"]))
    L.append("\\section*{PROJECTS}")
    for i, p in enumerate(d["projects"]):
        if i:
            L.append("\\vspace{1pt}")
        right = rf"\hfill \href{{{p['link'][1]}}}{{{tex_esc(p['link'][0])}}}" if p["link"] else ""
        L.append(rf"\noindent\textbf{{{tex_esc(p['title'])}}} ({tex_esc(p['stack'])}){right}\par")
        L.append(bullets(p["bullets"]))
    L.append("\\section*{TECHNICAL SKILLS}")
    for k, v in d["skills"]:
        L.append(rf"\textbf{{{tex_esc(k)}:}} {tex_esc(v)}\par")
    L.append("\n\\section*{ACHIEVEMENTS}")
    L.append(bullets(d["awards"]))
    L.append("\\end{document}\n")
    return "\n".join(L)


# ---------------------------------------------------------------- plain text
def build_txt(d) -> str:
    out = [d["name"].upper(), " | ".join(t for t, _ in d["contact"]), ""]
    def sec(n):
        out.extend([n.upper(), "-" * len(n)])
    sec("Summary"); out += [d["summary"], ""]
    sec("Education")
    for e in d["education"]:
        out.append(f"{e['title']}, {e['org']}" + (f" ({e['when']})" if e["when"] else ""))
        out += e["lines"]
    out.append("")
    sec("Experience")
    for e in d["experience"]:
        out.append(f"{e['title']}, {e['org']}" + (f" ({e['when']})" if e["when"] else ""))
        out += [f"- {b}" for b in e["bullets"]]
        out.append("")
    sec("Projects")
    for p in d["projects"]:
        out.append(f"{p['title']} ({p['stack']})" + (f" - {p['link'][0]}" if p["link"] else ""))
        out += [f"- {b}" for b in p["bullets"]]
        out.append("")
    sec("Technical Skills")
    out += [f"{k}: {v}" for k, v in d["skills"]]
    out.append("")
    sec("Achievements")
    out += [f"- {a}" for a in d["awards"]]
    return "\n".join(out) + "\n"


# ---------------------------------------------------------------- Word
def build_docx(d, path: Path):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(0.65)
    sec.top_margin = sec.bottom_margin = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    for sn in ("Heading 1", "List Bullet"):
        doc.styles[sn].paragraph_format.line_spacing = 1.0
    for sname in ("Heading 1", "List Bullet"):
        st = doc.styles[sname]
        st.font.name = "Arial"
        st.font.color.rgb = RGBColor(0, 0, 0)
    h1 = doc.styles["Heading 1"]
    h1.font.size = Pt(11)
    h1.font.bold = True
    h1.paragraph_format.space_before = Pt(6)
    h1.paragraph_format.space_after = Pt(2)
    doc.styles["List Bullet"].font.size = Pt(9.5)
    doc.styles["List Bullet"].paragraph_format.space_after = Pt(0)

    def heading(text):
        p = doc.add_paragraph(text.upper(), style="Heading 1")
        pPr = p._p.get_or_add_pPr()
        b = OxmlElement("w:pBdr")
        bt = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "000000")):
            bt.set(qn(k), v)
        b.append(bt)
        pPr.append(b)

    def line(bold, rest="", right="", space_before=0):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(space_before)
        # right-aligned tab stop at the text width, standard in ATS-safe resumes
        p.paragraph_format.tab_stops.add_tab_stop(Inches(7.2), alignment=2)
        r = p.add_run(bold)
        r.bold = True
        if rest:
            p.add_run(rest)
        if right:
            p.add_run("\t" + right)
        return p

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(d["name"])
    r.bold = True
    r.font.size = Pt(18)
    p = doc.add_paragraph(" | ".join(t for t, _ in d["contact"]))
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    heading("Summary")
    doc.add_paragraph(d["summary"])
    heading("Education")
    for e in d["education"]:
        line(e["title"], ", " + e["org"], e["when"])
        for ln in e["lines"]:
            doc.add_paragraph(ln)
    heading("Experience")
    for i, e in enumerate(d["experience"]):
        line(e["title"], ", " + e["org"], e["when"], 2 if i else 0)
        for b in e["bullets"]:
            doc.add_paragraph(b, style="List Bullet")
    heading("Projects")
    for i, pr in enumerate(d["projects"]):
        line(pr["title"], f" ({pr['stack']})", pr["link"][0] if pr["link"] else "", 2 if i else 0)
        for b in pr["bullets"]:
            doc.add_paragraph(b, style="List Bullet")
    heading("Technical Skills")
    for k, v in d["skills"]:
        p = doc.add_paragraph()
        p.add_run(k + ": ").bold = True
        p.add_run(v)
    heading("Achievements")
    for a in d["awards"]:
        doc.add_paragraph(a, style="List Bullet")
    doc.core_properties.author = d["name"]
    doc.core_properties.title = f"{d['name']} Resume"
    doc.save(path)


# ---------------------------------------------------------------- web page
def build_html(d) -> str:
    e = html.escape
    def a(t, u): return f'<a href="{e(u)}">{e(t)}</a>'
    body = [f"<h1>{e(d['name'])}</h1>",
            '<p class="contact">' + " &middot; ".join(a(t, u) for t, u in d["contact"]) + "</p>",
            "<h2>Summary</h2>", f"<p>{e(d['summary'])}</p>", "<h2>Education</h2>"]
    for x in d["education"]:
        body.append(f'<p class="row"><span><strong>{e(x["title"])}</strong>, {e(x["org"])}</span><span>{e(x["when"])}</span></p>')
        body += [f"<p>{e(l)}</p>" for l in x["lines"]]
    body.append("<h2>Experience</h2>")
    for x in d["experience"]:
        body.append(f'<p class="row"><span><strong>{e(x["title"])}</strong>, {e(x["org"])}</span><span>{e(x["when"])}</span></p>')
        if x["bullets"]:
            body.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in x["bullets"]) + "</ul>")
    body.append("<h2>Projects</h2>")
    for p in d["projects"]:
        link = a(*p["link"]) if p["link"] else ""
        body.append(f'<p class="row"><span><strong>{e(p["title"])}</strong> ({e(p["stack"])})</span><span>{link}</span></p>')
        body.append("<ul>" + "".join(f"<li>{e(b)}</li>" for b in p["bullets"]) + "</ul>")
    body.append("<h2>Technical Skills</h2>")
    body += [f"<p><strong>{e(k)}:</strong> {e(v)}</p>" for k, v in d["skills"]]
    body.append("<h2>Achievements</h2><ul>" + "".join(f"<li>{e(x)}</li>" for x in d["awards"]) + "</ul>")
    paper = "\n".join(body)

    dl = [
        ("PDF", f"{BASE}.pdf", "Best for applying. Text is selectable and parses cleanly."),
        ("Word (.docx)", f"{BASE}.docx", "For portals that ask for Word, or to edit."),
        ("LaTeX (.tex)", f"{BASE}.tex", "Source file. Compile with pdflatex."),
        ("Plain text (.txt)", f"{BASE}.txt", "For pasting into application forms."),
    ]
    buttons = "\n".join(
        f'<a class="btn{" primary" if i == 0 else ""}" href="{f}" download><strong>{e(n)}</strong><span>{e(h)}</span></a>'
        for i, (n, f, h) in enumerate(dl)
    )
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>Resume | Piyush Chaudhary</title>
<meta name="description" content="Resume of Piyush Chaudhary, Computer Engineering student and software developer. Download as PDF, Word, LaTeX or plain text.">
<link rel="canonical" href="https://piyushchy.github.io/resume/">
<link rel="icon" href="/favicon.svg" type="image/svg+xml">
<meta name="theme-color" content="#0a1317">
<style>
:root{{--ink:#0a1317;--ink2:#101c22;--paper:#e6eeee;--dim:#9db0b3;--sig:#7fd6e6;--line:rgba(230,238,238,.14)}}
*{{box-sizing:border-box}}
body{{margin:0;background:var(--ink);color:var(--paper);font:16px/1.55 system-ui,-apple-system,"Segoe UI",Roboto,Arial,sans-serif}}
a{{color:inherit}}
.wrap{{max-width:60rem;margin:0 auto;padding:1.25rem 1rem 4rem}}
.top{{display:flex;justify-content:space-between;align-items:center;gap:1rem;padding:.5rem 0 2rem}}
.top a{{font:12px/1 ui-monospace,Menlo,Consolas,monospace;letter-spacing:.14em;text-transform:uppercase;text-decoration:none;color:var(--dim)}}
.top a:hover,.top a:focus-visible{{color:var(--sig)}}
h1.title{{font-size:clamp(2.4rem,8vw,4.5rem);line-height:1;margin:0 0 .6rem;letter-spacing:-.02em}}
.lead{{color:var(--dim);max-width:38rem;margin:0 0 2rem}}
.dl{{display:grid;gap:.75rem;grid-template-columns:repeat(auto-fit,minmax(13rem,1fr));margin-bottom:.75rem}}
.btn{{display:flex;flex-direction:column;gap:.25rem;padding:1rem 1.1rem;border:1px solid var(--line);border-radius:.9rem;background:var(--ink2);text-decoration:none;transition:border-color .2s,transform .2s}}
.btn:hover,.btn:focus-visible{{border-color:var(--sig);transform:translateY(-2px)}}
.btn strong{{font-size:1.05rem}}
.btn span{{color:var(--dim);font-size:.85rem}}
.btn.primary{{background:var(--sig);color:#04181d;border-color:var(--sig)}}
.btn.primary span{{color:#0b3a44}}
.note{{color:var(--dim);font-size:.85rem;margin:0 0 2.5rem}}
.paper{{background:#f7f7f4;color:#111;border-radius:.75rem;padding:clamp(1.25rem,4vw,2.75rem);font:15px/1.5 Arial,Helvetica,sans-serif;box-shadow:0 30px 80px -30px rgba(0,0,0,.7)}}
.paper h1{{text-align:center;font-size:1.9rem;margin:0 0 .25rem}}
.paper .contact{{text-align:center;margin:0 0 .5rem;font-size:.9rem;overflow-wrap:anywhere}}
.paper h2{{font-size:1rem;text-transform:uppercase;letter-spacing:.04em;border-bottom:1px solid #111;margin:1.1rem 0 .4rem;padding-bottom:.15rem}}
.paper p{{margin:.15rem 0}}
.paper ul{{margin:.2rem 0 .6rem;padding-left:1.25rem}}
.paper li{{margin:.1rem 0}}
.paper .row{{display:flex;justify-content:space-between;gap:1rem;flex-wrap:wrap}}
.paper a{{color:#0b4f8a}}
@media print{{body{{background:#fff}}.top,.lead,.dl,.note,.title{{display:none}}.paper{{box-shadow:none;padding:0}}}}
</style>
</head>
<body>
<div class="wrap">
<nav class="top" aria-label="Site"><a href="/">&larr; Back to portfolio</a><a href="mailto:piyush.chaudhary@somaiya.edu">Email me</a></nav>
<main>
<h1 class="title">Resume</h1>
<p class="lead">Plain on purpose: one column, standard headings and real text, so applicant tracking systems read it correctly. Pick the format a portal asks for.</p>
<div class="dl">
{buttons}
</div>
<p class="note">All four files contain the same content.</p>
<article class="paper" aria-label="Resume">
{paper}
</article>
</main>
</div>
</body>
</html>
"""


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    tex = build_tex(DATA)
    (OUT / f"{BASE}.tex").write_text(tex, encoding="utf-8")
    (OUT / f"{BASE}.txt").write_text(build_txt(DATA), encoding="utf-8")
    build_docx(DATA, OUT / f"{BASE}.docx")
    (OUT / "index.html").write_text(build_html(DATA), encoding="utf-8")
    with tempfile.TemporaryDirectory() as td:
        shutil.copy(OUT / f"{BASE}.tex", td)
        for _ in range(2):
            r = subprocess.run(["pdflatex", "-interaction=nonstopmode", "-halt-on-error", f"{BASE}.tex"],
                               cwd=td, capture_output=True, text=True)
            if r.returncode:
                print(r.stdout[-2500:])
                raise SystemExit("pdflatex failed")
        shutil.copy(Path(td) / f"{BASE}.pdf", OUT / f"{BASE}.pdf")
    pages = re.search(r"Output written.*\((\d+) page", r.stdout)
    print("built", ", ".join(sorted(p.name for p in OUT.iterdir())), "| pages:", pages.group(1) if pages else "?")


if __name__ == "__main__":
    main()
